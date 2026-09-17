from pathlib import Path
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

OUT = Path(__file__).parent / 'resources' / 'template.docx'

TITLE = '施工中場所火災－消防安全管理及應變執行情形表'
NOTE = '備註：本表供火災案件初期彙整及首長備詢參考，實際內容應依現場調查及相關紀錄更新。'

ROWS = [
    ('項次', '項目', '執行情形'),
    ('SECTION', '一、基本概況', ''),
    ('1', '場所名稱', '○○○○'),
    ('2', '場所地址', '臺北市○○區○○路○段○號'),
    ('3', '火災時間', '115年○月○日○時○分'),
    ('4', '火災位置', '○樓○○區域'),
    ('5', '原因分析', '初步研判為○○○○；實際起火原因由火災調查單位調查中。'),
    ('SECTION', '二、消防安全管理制度執行情形', ''),
    ('1', '消防安全設備', '最近一次於○年○月○日實施消防安全檢查，檢查結果○○。'),
    ('2', '消防安全設備檢修申報', '最近一次於○年○月○日辦理○年○半年消防安全設備檢修申報。'),
    ('3', '防火管理', '消防防護計畫書為○年○月○日提報；另○年○半年自衛消防編組訓練於○年○月○日辦理。'),
    ('4', '防焰物品', '依法應使用防焰物品之區域，其窗簾、布幕、地毯等防焰物品情形：○○○○。'),
    ('SECTION', '三、施工中消防安全管理', ''),
    ('1', '施工中消防防護計畫', '於○年○月○日訂定（提報）施工中消防防護計畫。'),
    ('2', '消防安全設備停用情形', '施工期間消防安全設備：□無停用　□有停用；停用設備為：○○○○。'),
    ('3', '替代安全措施', '如有消防安全設備停用，替代安全措施為：○○○○。　□有　□無　□不適用'),
    ('SECTION', '四、火災當日應變情形', ''),
    ('1', '消防安全設備動作情形', '火災發生時，○○設備有／無動作；火警自動警報設備於○時○分發報，○○設備正常啟動。'),
    ('2', '119通報及初期滅火', '現場人員於○時○分通報119，並使用滅火器或消防設備進行初期滅火。'),
    ('3', '人員疏散及避難引導', '火災發生後，由自衛消防編組人員引導○○人依避難逃生動線疏散至○○處，並派員引導消防人員進入火災現場。'),
]


def set_run_font(run, size_pt=9.5, bold=None):
    run.font.name = '標楷體'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '標楷體')
    run.font.size = Pt(size_pt)
    run.bold = bold


def build_template(path=OUT):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Emu(7772400)
    sec.page_height = Emu(10058400)
    sec.top_margin = Emu(431800)
    sec.bottom_margin = Emu(431800)
    sec.left_margin = Emu(467995)
    sec.right_margin = Emu(467995)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    set_run_font(r, 16, True)

    table = doc.add_table(rows=len(ROWS), cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [431800, 1583690, 4824095]

    for i, rowdef in enumerate(ROWS):
        if rowdef[0] == 'SECTION':
            cell = table.rows[i].cells[0].merge(table.rows[i].cells[2])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            r = p.add_run(rowdef[1])
            set_run_font(r, 9.5, True)
            continue

        for j, txt in enumerate(rowdef):
            c = table.rows[i].cells[j]
            c.width = Emu(widths[j])
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            if i == 0 or j == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt)
            size = 9.0 if (i == 18 and j == 2) else 9.5
            set_run_font(r, size, True if i == 0 else None)

    for j, w in enumerate(widths):
        table.columns[j].width = Emu(w)

    p = doc.add_paragraph()
    r = p.add_run(NOTE)
    set_run_font(r, 8.5, None)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


if __name__ == '__main__':
    print(build_template())
