from __future__ import annotations

import re
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

from parser import PlaceData


PENDING_MARK = "【待補】"


def resource_path(relative: str) -> Path:
    """支援一般執行與 PyInstaller --onefile。"""
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base / relative


def roc_year(year: int) -> int:
    return year - 1911 if year >= 1912 else year


def roc_date(dt: datetime) -> str:
    return f"{roc_year(dt.year)}年{dt.month}月{dt.day}日"


def _result_phrase(result: str) -> str:
    if result in ("合格", "符合"):
        return "符合規定"
    if result in ("不合格", "不符合"):
        return "不符合規定"
    return result


def _set_cell_text_preserve_format(cell, text: str):
    """只換文字，不改表格、段落與第一個 run 的格式。"""
    if not cell.paragraphs:
        p = cell.add_paragraph()
    else:
        p = cell.paragraphs[0]

    if p.runs:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(text)

    # 清掉多餘段落內的文字，但不移除段落，避免碰版面結構。
    for extra_p in cell.paragraphs[1:]:
        for run in extra_p.runs:
            run.text = ""


def _mark_cell_pending(cell):
    """在尚未填入資料的執行情形前加上紅色【待補】，原制式預留文字保留。"""
    original = cell.text.strip()
    if not original or original.startswith(PENDING_MARK):
        return

    if not cell.paragraphs:
        p = cell.add_paragraph()
    else:
        p = cell.paragraphs[0]

    if p.runs:
        base_run = p.runs[0]
        base_rpr = deepcopy(base_run._r.rPr) if base_run._r.rPr is not None else None
        for run in p.runs:
            run.text = ""
    else:
        base_run = p.add_run()
        base_rpr = None

    base_run.text = PENDING_MARK
    base_run.bold = True
    base_run.font.color.rgb = RGBColor(255, 0, 0)

    normal = p.add_run(original)
    if base_rpr is not None:
        if normal._r.rPr is not None:
            normal._r.remove(normal._r.rPr)
        normal._r.insert(0, deepcopy(base_rpr))

    for extra_p in cell.paragraphs[1:]:
        for run in extra_p.runs:
            run.text = ""


def _sanitize_filename(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]', "_", name).strip().rstrip(".")
    return name[:80] or "場所"


def build_prefill_text(data: PlaceData) -> dict[str, str]:
    fields: dict[str, str] = {}

    if data.display_name:
        fields["場所名稱"] = data.display_name
    if data.address:
        fields["場所地址"] = data.address

    if data.latest_equipment_inspection:
        r = data.latest_equipment_inspection
        fields["消防安全設備"] = (
            f"最近一次於{roc_date(r.date)}實施消防安全檢查，"
            f"檢查結果{_result_phrase(r.result)}。"
        )

    if data.latest_equipment_report:
        r = data.latest_equipment_report
        report_year = roc_year(r.year)
        fields["消防安全設備檢修申報"] = (
            f"最近一次於{roc_date(r.received_date)}辦理"
            f"{report_year}年{r.period}消防安全設備檢修申報。"
        )

    if data.fire_plan_date:
        text = f"消防防護計畫書製定（變更）日期為{roc_date(data.fire_plan_date)}"
        if data.training_date:
            if data.training_year:
                year = roc_year(data.training_year)
                period = data.training_period or ""
                text += f"；另{year}年{period}自衛消防編組訓練於{roc_date(data.training_date)}辦理"
            else:
                text += f"；另自衛消防編組訓練於{roc_date(data.training_date)}辦理"
        fields["防火管理"] = text + "。"

    if data.flame_retardant_text:
        content = data.flame_retardant_text.rstrip("。")
        fields["防焰物品"] = (
            "依法應使用防焰物品之區域，其窗簾、布幕、地毯等防焰物品情形："
            f"{content}。"
        )

    return fields


def fill_word_template(
    data: PlaceData,
    output_path: str | Path,
    template_path: str | Path | None = None,
) -> tuple[Path, list[str]]:
    template_path = Path(template_path) if template_path else resource_path("resources/template.docx")
    output_path = Path(output_path)

    doc = Document(template_path)
    if not doc.tables:
        raise ValueError("Word 樣板內找不到表格。")
    table = doc.tables[0]

    fields = build_prefill_text(data)
    updated: list[str] = []

    # 不依列號猜欄位：直接找第二欄固定「項目」名稱，再只改第三欄。
    for row in table.rows:
        cells = row.cells
        if len(cells) < 3:
            continue
        item = cells[1].text.strip()
        if not item or item == "項目" or item.startswith(("一、", "二、", "三、", "四、")):
            continue

        if item in fields:
            _set_cell_text_preserve_format(cells[2], fields[item])
            updated.append(item)
        else:
            _mark_cell_pending(cells[2])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path, updated


def default_output_path(pdf_path: str | Path, data: PlaceData) -> Path:
    pdf_path = Path(pdf_path)
    stem = _sanitize_filename(data.sign_name or data.company_name or data.raw_place_name)
    return pdf_path.with_name(f"{stem}_施工中場所火災_消防安全管理及應變執行情形表.docx")
